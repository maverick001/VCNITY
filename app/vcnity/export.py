"""Markdown → .docx, enough for a report: headings, paragraphs, bullets, one
pipe table rendered as a real Word table."""
from __future__ import annotations

import re
from pathlib import Path


def _add_runs(paragraph, text: str) -> None:
    # **bold** and _italic_ only — that's all the report uses
    for part in re.split(r"(\*\*[^*]+\*\*|_[^_]+_)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("_") and part.endswith("_") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def report_to_docx(markdown: str, path: Path) -> Path:
    import docx

    d = docx.Document()
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                t = d.add_table(rows=len(rows), cols=len(rows[0]))
                t.style = "Table Grid"
                for r, cells in enumerate(rows):
                    for c, val in enumerate(cells[: len(rows[0])]):
                        cell = t.cell(r, c)
                        cell.text = ""
                        _add_runs(cell.paragraphs[0], val)
                        if r == 0:
                            for run in cell.paragraphs[0].runs:
                                run.bold = True
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            d.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
        elif line.lstrip().startswith(("- ", "* ")):
            _add_runs(d.add_paragraph(style="List Bullet"), line.lstrip()[2:])
        else:
            _add_runs(d.add_paragraph(), line.strip())
        i += 1
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(path))
    return path
