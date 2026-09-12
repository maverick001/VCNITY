import io
from pathlib import Path

import numpy as np
import pytest

from vcnity.models import SourceFile
from vcnity.stages import s0_intake, s1_ingest


def _docx(path: Path):
    import docx

    d = docx.Document()
    d.add_paragraph("First paragraph about the carpark.")
    d.add_paragraph("Second paragraph about lighting.")
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "cell one"
    t.rows[0].cells[1].text = "cell two"
    d.save(path)


def _png_bytes() -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), (200, 30, 30)).save(buf, format="PNG")
    return buf.getvalue()


def _pptx(path: Path, tmp: Path):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Slide title here"
    tb = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    tb.text_frame.text = "A text box on the slide"
    png = tmp / "pic.png"
    png.write_bytes(_png_bytes())
    slide.shapes.add_picture(str(png), Inches(1), Inches(3))
    slide.notes_slide.notes_text_frame.text = "speaker notes"
    prs.save(path)


def _wav(path: Path, sr=48000, seconds=1.0, channels=2):
    import av

    n = int(sr * seconds)
    t = np.arange(n) / sr
    tone = (0.3 * np.sin(2 * np.pi * 440 * t)).astype(np.float32)
    data = np.stack([tone] * channels)
    with av.open(str(path), "w") as out:
        stream = out.add_stream("pcm_s16le", rate=sr)
        stream.layout = "stereo" if channels == 2 else "mono"
        frame = av.AudioFrame.from_ndarray(
            (data * 32767).astype(np.int16).T.reshape(1, -1) if channels == 2 else (data * 32767).astype(np.int16),
            format="s16", layout=stream.layout.name,
        )
        frame.sample_rate = sr
        for p in stream.encode(frame):
            out.mux(p)
        for p in stream.encode(None):
            out.mux(p)


def test_extract_docx(tmp_path):
    p = tmp_path / "a.docx"
    _docx(p)
    text, images = s1_ingest.extract_text(p)
    assert "First paragraph" in text and "cell two" in text
    assert images == []


def test_extract_pptx_with_picture(tmp_path):
    p = tmp_path / "a.pptx"
    _pptx(p, tmp_path)
    text, images = s1_ingest.extract_text(p, out_dir=tmp_path / "out")
    assert "Slide title here" in text and "text box" in text and "speaker notes" in text
    assert len(images) == 1 and images[0].exists()


def test_extract_txt(tmp_path):
    p = tmp_path / "a.txt"
    p.write_text("plain words", encoding="utf-8")
    assert s1_ingest.extract_text(p)[0] == "plain words"


def test_audio_to_16k_mono(tmp_path):
    src = tmp_path / "in.wav"
    _wav(src)
    out = s1_ingest.audio_to_wav16k(src, tmp_path / "out.wav")
    import av

    with av.open(str(out)) as c:
        s = c.streams.audio[0]
        assert s.rate == 16000 and s.channels == 1
        assert 0.9 < float(c.duration / 1e6) < 1.2


def test_intake_levels(db_session, tmp_path):
    job = s0_intake.create_job(db_session, "j")
    f = tmp_path / "x.txt"
    f.write_text("hi")
    sf = s0_intake.add_file(db_session, job.id, f, level=2, consent_label="c1", consent_scope="session")
    assert sf.kind == "text" and sf.level == 2 and sf.consent_id is not None
    with pytest.raises(ValueError):
        s0_intake.set_level(db_session, sf.id, 1, actor_role="analyst")
    s0_intake.set_level(db_session, sf.id, 3, actor_role="community")
    assert db_session.get(SourceFile, sf.id).level == 3
    s0_intake.confirm_level(db_session, sf.id)
    assert db_session.get(SourceFile, sf.id).level_confirmed_by_community is True


def test_intake_kinds(db_session, tmp_path):
    job = s0_intake.create_job(db_session, "j")
    for name, kind in [("a.m4a", "audio"), ("b.jpeg", "image"), ("c.pptx", "text"), ("d.xlsx", "brief")]:
        p = tmp_path / name
        p.write_bytes(b"x")
        assert s0_intake.add_file(db_session, job.id, p, level=1, consent_label="c", consent_scope="s").kind == kind


def test_ingest_brief_and_text(db_session, tmp_path):
    import openpyxl

    job = s0_intake.create_job(db_session, "j")
    xl = tmp_path / "b.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Project Link", "Background"])
    ws.append(["P1", "The community was invited to help."])
    ws.append(["P2", "Residents shared experiences."])
    wb.save(xl)
    s0_intake.add_file(db_session, job.id, xl, level=1, consent_label="client", consent_scope="brief")
    t = tmp_path / "n.txt"
    t.write_text("Facilitator notes about the session.")
    s0_intake.add_file(db_session, job.id, t, level=2, consent_label="c", consent_scope="s")
    results = s1_ingest.run(db_session, job.id, cache_dir=tmp_path / "cache")
    db_session.refresh(job)
    assert "invited to help" in job.brief and "shared experiences" in job.brief
    txt = [r for r in results if r["kind"] == "text"][0]
    assert txt["chars"] > 10 and Path(txt["ingested_path"]).exists()
